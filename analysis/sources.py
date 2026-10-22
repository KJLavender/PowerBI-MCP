"""Read the report's source data the way the Power BI model does (from the M partitions).

  Oracle.Database(...)                 → the Oracle connection of the simulator (.env)
  Excel.Workbook(File.Contents(path))  → the workbook's sheets, loaded into an in-memory DuckDB

Both give a DB-API style connection that accepts the same SQL (named parameters :d0 / :d1),
so snapshots are taken with one code path whatever the source.
"""
import decimal
import re
from datetime import date, timedelta
from pathlib import Path

import duckdb
import pandas as pd

from . import dataset
from .context import Context
from .semantic import cond_dax, dax_col, dax_table, ident

BLANK = "(空白)"  # same label as store.BLANK


def excel_path(ref: str) -> Path:
    """The path in the model is absolute on the machine that built it; fall back to the dataset folder."""
    p = Path(ref)
    if p.exists():
        return p
    local = dataset.HOME / "data" / Path(ref.replace("\\", "/")).name
    if local.exists():
        return local
    raise FileNotFoundError(f"Excel source {ref} not found (also looked for {local})")


class SourceError(SystemExit):
    """The files alone cannot reproduce the report's data."""


def read_file_table(t, ctx: Context) -> pd.DataFrame:
    """One Excel sheet / Excel table / CSV file, read the way its Power Query does."""
    src = t.source
    if src.get("steps"):
        raise SourceError(
            f"資料表 {t.name} 的 Power Query 有會改變資料列的步驟（{'、'.join(src['steps'])}），直接讀檔案會和 Power BI "
            f"不一致。請在分析設定加上 source: {{kind: powerbi, title: <報表視窗標題>}}，改從開啟中的報表讀取。")
    path = excel_path(src["file"])
    if dataset.VARIANT:
        path = path.with_name(ctx.config["variants"][dataset.VARIANT])
    header = 0 if src.get("promote") else None
    if src["kind"] == "csv":
        df = pd.read_csv(path, sep=src["delimiter"], encoding=src["encoding"], header=header)
    elif src["item_kind"] == "Table":  # an Excel table object: its own range, header row included
        import openpyxl
        book = openpyxl.load_workbook(path, data_only=True, read_only=False)
        ws, ref = next(((ws, ws.tables[t.source_table].ref) for ws in book.worksheets
                        if t.source_table in ws.tables), (None, None))
        if ws is None:
            raise SourceError(f"{path.name} 裡找不到 Excel 表格 {t.source_table}")
        rows = [[c.value for c in r] for r in ws[ref]]
        df = pd.DataFrame(rows[1:], columns=rows[0])
    elif src["item_kind"] == "Sheet":
        df = pd.read_excel(path, sheet_name=t.source_table, header=header)
    else:
        raise SourceError(f"資料表 {t.name} 讀的是 Excel 的 {src['item_kind']}，請改用 source: powerbi")
    if header is None and not (src["kind"] == "excel" and src["item_kind"] == "Table"):
        # no header row promoted: Power Query names the columns Column1, Column2, … (Excel tables keep theirs)
        df.columns = [f"Column{i + 1}" for i in range(df.shape[1])]
    df = df.rename(columns={a.replace('""', '"'): b.replace('""', '"') for a, b in src.get("renames", [])})
    return apply_types(df, t.m)


def apply_types(df: pd.DataFrame, m: str) -> pd.DataFrame:
    """Table.TransformColumnTypes: text from a CSV becomes dates and numbers, as in Power BI."""
    call = m[m.find("Table.TransformColumnTypes"):] if "Table.TransformColumnTypes" in m else ""
    for col, typ in re.findall(r'\{\s*"((?:[^"]|"")*)"\s*,\s*([^}]+?)\s*\}', call):
        col = col.replace('""', '"')
        if col not in df.columns:
            continue
        typ = typ.lower()
        if "date" in typ:
            df[col] = pd.to_datetime(df[col], errors="coerce")
            if "datetime" not in typ:
                df[col] = df[col].dt.normalize()
        elif any(k in typ for k in ("int64", "number", "currency", "percentage", "double", "decimal")):
            df[col] = pd.to_numeric(df[col], errors="coerce")
        elif "text" in typ:
            df[col] = df[col].astype("string")
    return df


class ExcelSource:
    """Excel sheets / tables and CSV files as DuckDB tables, named like the model expects (store.py SQL)."""

    def __init__(self, ctx: Context):
        self.db = duckdb.connect()
        needed = {ctx.fact} | {d.table for d in ctx.dims.values()} | {ctx.date_dim.table} | \
            {t for b in ctx.bases.values() for t in b.tables} | {c[0] for c in ctx.report_filters}
        for t in ctx.model.tables.values():
            if t.source.get("kind") not in ("excel", "csv") or t.name not in needed:
                continue
            df = read_file_table(t, ctx)
            name = t.source_table or t.name
            self.db.register(f"_{t.name}", df)
            self.db.execute(f"CREATE TABLE {ident(name)} AS SELECT * FROM {ident('_' + t.name)}")
        self._res = None

    # --- minimal DB-API surface used by store.py ---
    def cursor(self):
        return self

    def execute(self, sql: str, **params):
        self._res = self.db.execute(re.sub(r"(?<!:):(\w+)", r"$\1", sql), params)
        return self

    def fetchall(self):
        return self._res.fetchall()

    def fetchone(self):
        return self._res.fetchone()

    @property
    def description(self):
        return self._res.description

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        self.db.close()


class PbixSource(ExcelSource):
    """The data stored inside a .pbix (what Power BI imported, after Power Query), read without Power BI: PBIXRay
    decodes the tables, which become DuckDB tables named like the model's sources so store.py's SQL runs as is."""

    def __init__(self, ctx: Context):
        self.db = duckdb.connect()
        self._res = None
        path = Path(ctx.source["file"])
        if not path.exists():
            path = dataset.HOME / path.name
        from pbixray import PBIXRay
        px = PBIXRay(str(path))
        for t in ctx.model.tables.values():
            if t.name in needed_tables(ctx):
                register_table(self.db, px.get_table(t.name), t.source_table or t.name,
                               {c.name: c.source_column for c in t.columns.values() if c.source_column})


def register_table(db, df: pd.DataFrame, name: str, renames: dict[str, str]) -> None:
    """A model table as a DuckDB table under its source name, columns under their source column names."""
    df = df.rename(columns={k: v for k, v in renames.items() if k in df.columns and k != v})
    for col in df.columns:  # fixed-decimal (currency) columns come as Decimal objects: as floats, like the snapshots
        first = df[col].dropna().head(1)
        if df[col].dtype == object and len(first) and isinstance(first.iloc[0], decimal.Decimal):
            df[col] = pd.to_numeric(df[col]).astype("float64")
    db.register("_pbix_table", df)
    db.execute(f"CREATE OR REPLACE TABLE {ident(name)} AS SELECT * FROM _pbix_table")
    db.unregister("_pbix_table")


def needed_tables(ctx: Context) -> set[str]:
    return {ctx.fact} | {d.table for d in ctx.dims.values()} | {ctx.date_dim.table} | \
        {t for b in ctx.bases.values() for t in b.tables} | {c[0] for c in ctx.report_filters} | \
        ({ctx.dq["etl_table"]} if ctx.dq.get("etl_table") else set())


def connect(ctx: Context):
    if ctx.source["kind"] == "pbix":
        return PbixSource(ctx)
    if ctx.source["kind"] in ("excel", "csv"):
        return ExcelSource(ctx)
    from simulator.db import connect as oracle_connect
    return oracle_connect()


def date_range(ctx: Context) -> tuple[date, date]:
    """First and last business date available in the source."""
    if ctx.source["kind"] == "powerbi":
        return PowerBISource(ctx).date_range()
    if ctx.source["kind"] == "pbix":  # the days the fact table has rows (the date table may span more years)
        from .store import _from_clause
        from_sql, alias = _from_clause(ctx, [ctx.date_dim])
        col = f"{alias[ctx.date_dim.table]}.{ident(ctx.date_dim.source_column)}"
        with connect(ctx) as src:
            lo, hi = src.execute(f"SELECT min({col}), max({col}) {from_sql}").fetchone()
        return pd.Timestamp(lo).date(), pd.Timestamp(hi).date()
    if ctx.source["kind"] in ("excel", "csv"):
        with connect(ctx) as src:
            t = ctx.model.tables[ctx.date_dim.table]
            col = ident(ctx.date_dim.source_column)
            lo, hi = src.execute(f"SELECT min({col}), max({col}) FROM {ident(t.source_table or t.name)}").fetchone()
        return pd.Timestamp(lo).date(), pd.Timestamp(hi).date()
    from simulator import config as sim_cfg
    return sim_cfg.HISTORY_START, date.today() - timedelta(days=1)


def report_filter_dax(ctx: Context) -> list[str]:
    """Report-level filters as DAX filter tables: FILTER(ALL('T'[c]), condition)."""
    return [f"FILTER(ALL({dax_col(c[0], c[1])}), {cond_dax(c)})" for c in ctx.report_filters]


class PowerBISource:
    """The data already imported into an open Power BI Desktop report, read with DAX.

    Works whatever the report's own source is (SQL Server, SharePoint, CSV …): the engine answers the same
    grouped aggregates the SQL path asks the database for. Blank members come back as BLANK → "(空白)"."""

    def __init__(self, ctx: Context):
        from . import pbi
        self.ctx, self.pbi = ctx, pbi
        self.port = pbi.find_port(ctx.source.get("title"))
        self.date = dax_col(ctx.date_dim.table, ctx.date_dim.column)

    def between(self, d0: date, d1: date) -> str:
        """The date range plus the report-level filters, as SUMMARIZECOLUMNS filter arguments."""
        return ", ".join([f"FILTER(ALL({self.date}), {self.date} >= DATE({d0.year},{d0.month},{d0.day}) && "
                          f"{self.date} <= DATE({d1.year},{d1.month},{d1.day}))"] + report_filter_dax(self.ctx))

    def query(self, columns: list[str], names: list[str], d0: date, d1: date, measures: dict[str, str]) -> pd.DataFrame:
        exprs = "".join(f', "{k}", {v}' for k, v in measures.items())
        q = f"EVALUATE SUMMARIZECOLUMNS({', '.join(columns + [self.between(d0, d1)])}{exprs})"
        (df,) = self.pbi.frames(self.port, [q])
        df.columns = names + list(measures)
        return df

    def grouped(self, dim_keys: list[str], bases: list[str], d0: date, d1: date) -> pd.DataFrame:
        dims = [self.ctx.dims[k] for k in dim_keys]
        df = self.query([self.date] + [dax_col(d.table, d.column) for d in dims], ["business_date"] + dim_keys,
                        d0, d1, {b: self.ctx.bases[b].dax for b in bases})
        for k in dim_keys:
            df[k] = df[k].fillna(BLANK)
        for b in bases:
            df[b] = pd.to_numeric(df[b]).astype(float).fillna(0)  # float: an all-integer first month must not make the column integer
        if not df.empty:
            df["business_date"] = pd.to_datetime(df["business_date"]).dt.date
        return df

    def order_level_dims(self, d1: date) -> tuple[list[str], set[str]]:
        """Dimensions with one value per distinct-count key (e.g. ORDER_ID), checked on the last 14 days; and the
        keys that span several days (their daily distinct counts cannot be added up)."""
        keys = {b.distinct_of for b in self.ctx.bases.values() if b.distinct_of}
        result, cross_day = set(self.ctx.dims), set()
        for key in keys:
            col = next(c for c in self.ctx.model.tables[self.ctx.fact].columns.values()
                       if (c.source_column or c.name) == key)
            dims = list(self.ctx.dims)
            df = self.query([dax_col(self.ctx.fact, col.name), self.date] +
                            [dax_col(self.ctx.dims[k].table, self.ctx.dims[k].column) for k in dims],
                            ["key", "business_date"] + dims, d1 - timedelta(days=13), d1,
                            {"n": self.ctx.bases["rows"].dax})
            if df.groupby("key")["business_date"].nunique().max() > 1:
                cross_day.add(key)
                continue
            per_key = df.groupby("key")[dims].nunique(dropna=False).max()
            result &= {k for k in dims if per_key[k] <= 1}
        return [k for k in self.ctx.dims if k in result], cross_day

    def calendar(self, d0: date, d1: date) -> pd.DataFrame:
        if self.ctx.holiday_source:
            holiday = next(c.name for c in self.ctx.model.tables[self.ctx.date_dim.table].columns.values()
                           if c.source_column == self.ctx.holiday_source)
            df = self.query([self.date, dax_col(self.ctx.date_dim.table, holiday)], ["business_date", "holiday"],
                            d0, d1, {})
            df["business_date"] = pd.to_datetime(df["business_date"]).dt.date
            return df.drop_duplicates("business_date")
        days = [d0 + timedelta(days=i) for i in range((d1 - d0).days + 1)]
        return pd.DataFrame({"business_date": days, "holiday": [None] * len(days)})

    def etl(self, d0: date, d1: date) -> pd.DataFrame | None:
        dq = self.ctx.dq
        if not dq.get("etl_table"):
            return None
        t = dq["etl_table"]
        df = self.query([self.date, dax_col(t, dq["etl_rows"]), dax_col(t, dq["etl_status"])],
                        ["business_date", "etl_rows", "etl_status"], d0, d1, {"n": f"COUNTROWS({dax_table(t)})"})
        df = df.drop(columns="n")
        df["etl_rows"] = pd.to_numeric(df["etl_rows"])
        if not df.empty:
            df["business_date"] = pd.to_datetime(df["business_date"]).dt.date
        return df

    def date_range(self) -> tuple[date, date]:
        has = f"NOT ISBLANK(CALCULATE({self.ctx.bases['rows'].dax}))"
        q = (f'EVALUATE ROW("a", FORMAT(MINX(FILTER(ALL({self.date}), {has}), {self.date}), "yyyy-MM-dd"), '
             f'"b", FORMAT(MAXX(FILTER(ALL({self.date}), {has}), {self.date}), "yyyy-MM-dd"))')
        (a, b), = self.pbi.run_dax(self.port, [q])[0]["rows"]
        return date.fromisoformat(a), date.fromisoformat(b)
